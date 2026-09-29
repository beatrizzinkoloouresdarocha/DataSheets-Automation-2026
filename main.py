import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

# ==============================================================================
# 1. CONFIGURAÇÃO DA URL DO GOOGLE PLANILHAS
# Link de exportação direta CSV da sua planilha
# ==============================================================================
URL_SHEET_CSV = "https://docs.google.com/spreadsheets/d/1a7odnKOYZziF_kIBnTIme0oGeCzkmFL1B3rmUclEWKA/export?format=csv"

print("Baixando e processando dados do Google Planilhas...")

try:
    # Tenta carregar os dados diretamente da URL do Google Sheets publicado
    df = pd.read_csv(URL_SHEET_CSV)
    print("✓ Dados carregados com sucesso a partir da URL pública!")
except (ValueError, OSError, RuntimeError) as e:
    print(f"⚠️ Não foi possível carregar da URL informada ({e}).")
    print("👉 Utilizando base de dados fictícia local com os 20 registros...")

    # Dados de fallback idênticos aos cadastrados na planilha online
    data_mock = {
        "ID": list(range(101, 121)),
        "Data": [
            "2026-01-05",
            "2026-01-06",
            "2026-01-08",
            "2026-01-10",
            "2026-01-12",
            "2026-01-15",
            "2026-01-18",
            "2026-01-20",
            "2026-01-22",
            "2026-01-25",
            "2026-02-01",
            "2026-02-03",
            "2026-02-05",
            "2026-02-08",
            "2026-02-10",
            "2026-02-12",
            "2026-02-15",
            "2026-02-18",
            "2026-02-20",
            "2026-02-22",
        ],
        "Categoria": [
            "Eletrônicos",
            "Roupas",
            "Eletrônicos",
            "Alimentos",
            "Roupas",
            "Eletrônicos",
            "Alimentos",
            "Roupas",
            "Eletrônicos",
            "Alimentos",
            "Eletrônicos",
            "Roupas",
            "Alimentos",
            "Eletrônicos",
            "Roupas",
            "Alimentos",
            "Eletrônicos",
            "Roupas",
            "Alimentos",
            "Eletrônicos",
        ],
        "Produto": [
            "Mouse Gamer",
            "Camiseta Algodão",
            "Teclado Mecânico",
            "Café Gourmet",
            "Calça Jeans",
            "Monitor 24",
            "Azeite Extra Virgem",
            "Jaqueta Couro",
            "Fone Bluetooth",
            "Chocolate Amargo",
            "Webcam HD",
            "Tênis Esportivo",
            "Vinho Tinto",
            "Carregador Rápido",
            "Meias Kit",
            "Queijo Parmesão",
            "Cadeira Gamer",
            "Boné Aba Reta",
            "Castanha de Caju",
            "Mousepad Extra Grande",
        ],
        "Quantidade": [
            5,
            12,
            3,
            20,
            8,
            2,
            15,
            4,
            10,
            30,
            6,
            5,
            8,
            15,
            25,
            10,
            1,
            14,
            18,
            8,
        ],
        "Vendas": [
            450,
            600,
            900,
            500,
            1200,
            1800,
            675,
            1600,
            1500,
            450,
            1200,
            1250,
            960,
            750,
            375,
            600,
            1500,
            560,
            720,
            320,
        ],
    }
    df = pd.DataFrame(data_mock)

# ==============================================================================
# 2. PROCESSAMENTO E ANÁLISE DOS DADOS
# ==============================================================================
df["Vendas"] = pd.to_numeric(df["Vendas"], errors="coerce").fillna(0)
df["Quantidade"] = pd.to_numeric(df["Quantidade"], errors="coerce").fillna(0)

resumo_categoria = (
    df.groupby("Categoria")[["Quantidade", "Vendas"]].sum().reset_index()
)

total_vendas = float(df["Vendas"].sum())
total_qtd = int(df["Quantidade"].sum())

# ==============================================================================
# 3. CRIAÇÃO DO RELATÓRIO FORMATADO EM EXCEL (OPENPYXL)
# ==============================================================================
wb = Workbook()

ws = wb.active
assert isinstance(ws, Worksheet)

ws.title = "Resumo Executivo"
ws.views.sheetView[0].showGridLines = True

COR_CABECALHO_PRINCIPAL = "1F4E79"
COR_CABECALHO_TABELA = "2F5597"
COR_LINHA_TOTAL = "D9E1F2"

font_titulo = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
font_cabecalho = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
font_bold = Font(name="Calibri", size=11, bold=True)
font_normal = Font(name="Calibri", size=11)

fill_titulo = PatternFill(
    start_color=COR_CABECALHO_PRINCIPAL, fill_type="solid"
)
fill_cabecalho = PatternFill(
    start_color=COR_CABECALHO_TABELA, fill_type="solid"
)
fill_total = PatternFill(start_color=COR_LINHA_TOTAL, fill_type="solid")

borda_fina = Side(style="thin", color="BFBFBF")
borda_dupla = Side(style="double", color="000000")
border_cell = Border(
    left=borda_fina, right=borda_fina, top=borda_fina, bottom=borda_fina
)
border_total = Border(top=borda_fina, bottom=borda_dupla)

ws.merge_cells("A1:C1")
cell_a1 = ws["A1"]
assert cell_a1 is not None
cell_a1.value = "DASHBOARD GERENCIAL DE VENDAS"
cell_a1.font = font_titulo
cell_a1.fill = fill_titulo
cell_a1.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 40

headers = ["Categoria", "Quantidade Total", "Total Vendas (R$)"]
for col_num, header in enumerate(headers, start=1):
    cell = ws.cell(row=3, column=col_num, value=header)
    cell.font = font_cabecalho
    cell.fill = fill_cabecalho
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = border_cell
ws.row_dimensions[3].height = 25

linha_atual = 4
for _, row in resumo_categoria.iterrows():
    c1 = ws.cell(row=linha_atual, column=1, value=str(row["Categoria"]))
    c2 = ws.cell(row=linha_atual, column=2, value=int(row["Quantidade"]))
    c3 = ws.cell(row=linha_atual, column=3, value=float(row["Vendas"]))

    c1.alignment = Alignment(horizontal="left", vertical="center")
    c2.alignment = Alignment(horizontal="center", vertical="center")
    c3.alignment = Alignment(horizontal="right", vertical="center")

    c2.number_format = "#,##0"
    c3.number_format = "R$ #,##0.00"

    for c in (c1, c2, c3):
        c.font = font_normal
        c.border = border_cell
    linha_atual += 1

c1_tot = ws.cell(row=linha_atual, column=1, value="TOTAL GERAL")
c2_tot = ws.cell(row=linha_atual, column=2, value=total_qtd)
c3_tot = ws.cell(row=linha_atual, column=3, value=total_vendas)

c1_tot.alignment = Alignment(horizontal="left", vertical="center")
c2_tot.alignment = Alignment(horizontal="center", vertical="center")
c3_tot.alignment = Alignment(horizontal="right", vertical="center")

c2_tot.number_format = "#,##0"
c3_tot.number_format = "R$ #,##0.00"

for c in (c1_tot, c2_tot, c3_tot):
    c.font = font_bold
    c.fill = fill_total
    c.border = border_total

ws.column_dimensions["A"].width = 22
ws.column_dimensions["B"].width = 20
ws.column_dimensions["C"].width = 22

# ==============================================================================
# 4. CRIAR E ADICIONAR O GRÁFICO DE BARRAS
# ==============================================================================
chart = BarChart()
chart.type = "col"
chart.style = 10
chart.title = "Total de Vendas por Categoria"
chart.y_axis.title = "Faturamento (R$)"
chart.x_axis.title = "Categoria"
chart.width = 16
chart.height = 10

data = Reference(ws, min_col=3, min_row=3, max_row=linha_atual - 1)
cats = Reference(ws, min_col=1, min_row=4, max_row=linha_atual - 1)

chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)

ws.add_chart(chart, "E3") # type: ignore

# ==============================================================================
# 5. SALVAR O ARQUIVO FINAL
# ==============================================================================
nome_relatorio = "Relatorio_Final_Analisado.xlsx"
wb.save(nome_relatorio)

print("--------------------------------------------------")
print(f"✅ Sucesso! Relatório gerado: {nome_relatorio}")
print("--------------------------------------------------")